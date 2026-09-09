"""Топ вложенных запросов (GetTop): что именно ищут по нашим фразам.

Метод отдаёт срез за последние 30 дней, а не ряд: `results` — вложенные
запросы с частотностью, `associations` — похожие запросы («с этим ищут»).
Даты в ответе нет, поэтому снимок помечается датой прогона.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import pathlib

from .api import WordstatError, validate_phrase
from .build import region_order, sheet_title
from .fetch import slug

COLUMNS = [
    "fetched_at", "phrase", "region_name", "region_id", "region_type", "device",
    "kind", "rank", "nested_phrase", "count", "total_count",
]


def run_name(run_date: dt.date, tag: str = "") -> str:
    """Имя прогона. Метка разводит прогоны с разными наборами фраз и глубиной."""
    return f"{run_date.isoformat()}-top" + (f"-{tag}" if tag else "")


def run_dir(cfg, run_date: dt.date, tag: str = "") -> pathlib.Path:
    return cfg.raw_dir / run_name(run_date, tag)


def select_phrases(cfg, only=None) -> tuple[str, ...]:
    """Подмножество фраз из конфига. Опечатка в названии — ошибка, не тишина."""
    if not only:
        return tuple(cfg.phrases)
    known = {p.casefold(): p for p in cfg.phrases}
    chosen, unknown = [], []
    for name in only:
        key = str(name).strip().casefold()
        (chosen.append(known[key]) if key in known else unknown.append(name))
    if unknown:
        raise SystemExit(
            "нет таких фраз в config.yaml: " + ", ".join(map(str, unknown)) +
            ".\nЕсть: " + ", ".join(cfg.phrases)
        )
    return tuple(dict.fromkeys(chosen))


def request_body(cfg, phrase: str, region_id: str, device: str, limit: int) -> dict:
    return {"phrase": phrase, "regions": [region_id], "devices": [device], "numPhrases": limit}


def filename(body: dict) -> str:
    digest = hashlib.sha1(
        json.dumps(body, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()[:10]
    return f"{body['regions'][0]}_{slug(body['phrase'])}_{slug(body['devices'][0])}_{digest}.json"


def fetch(cfg, client, run_date: dt.date | None = None, limit: int = 50,
          only=None, tag: str = "", log=print) -> dict:
    """Снимок топа по парам «фраза × регион». Кэш — как в основной выгрузке."""
    run_date = run_date or dt.date.today()
    phrases = select_phrases(cfg, only)
    directory = run_dir(cfg, run_date, tag)
    directory.mkdir(parents=True, exist_ok=True)
    summary = {
        "run_date": run_date.isoformat(),
        "method": "GetTop",
        "limit": limit,
        "tag": tag,
        "phrases": list(phrases),
        "planned": 0, "fetched": 0, "from_cache": 0, "errors": [],
    }

    plan = [
        (phrase, region, device)
        for phrase in phrases
        for region in cfg.resolved_regions()
        for device in cfg.devices
    ]
    summary["planned"] = len(plan)

    for index, (phrase, region, device) in enumerate(plan, 1):
        body = request_body(cfg, phrase, region.region_id, device, limit)
        path = directory / filename(body)
        label = f"{region.name} · {phrase} · {device}"
        if path.exists():
            try:
                stored = json.loads(path.read_text(encoding="utf-8"))
                if stored.get("request") == body and "response" in stored:
                    summary["from_cache"] += 1
                    continue
            except (ValueError, OSError):
                pass
        try:
            validate_phrase(phrase)
            response = client.get_top(
                phrase, regions=[region.region_id], devices=[device], num_phrases=limit
            )
        except (WordstatError, ValueError) as exc:
            summary["errors"].append({"call": label, "error": str(exc)})
            log(f"  [{index}/{len(plan)}] ОШИБКА  {label}: {exc}")
            continue
        path.write_text(
            json.dumps(
                {
                    "meta": {
                        "fetched_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
                        "region_name": region.name,
                        "region_type": region.type,
                    },
                    "request": body,
                    "response": response,
                },
                ensure_ascii=False, indent=2,
            ),
            encoding="utf-8",
        )
        summary["fetched"] += 1
        log(f"  [{index}/{len(plan)}] ок      {label}")

    summary["http_calls"] = client.stats["http_calls"]
    summary["billable_calls"] = client.stats["billable_calls"]
    summary["cost_rub"] = client.cost_rub
    (directory / "manifest.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return summary


def to_rows(cfg, records: list[dict]) -> list[dict]:
    by_id = {r.region_id: r for r in cfg.regions if r.region_id}
    order = region_order(cfg)
    rows = []
    for record in records:
        body = record["request"]
        meta = record.get("meta") or {}
        response = record["response"]
        region_id = body["regions"][0]
        region = by_id.get(str(region_id))
        region_name = meta.get("region_name") or (region.name if region else str(region_id))
        total = response.get("totalCount")
        for kind, key in (("nested", "results"), ("association", "associations")):
            for rank, item in enumerate(response.get(key) or [], 1):
                rows.append(
                    {
                        "fetched_at": (meta.get("fetched_at") or "")[:10],
                        "phrase": body["phrase"],
                        "region_name": region_name,
                        "region_id": str(region_id),
                        "region_type": meta.get("region_type") or (region.type if region else ""),
                        "device": body["devices"][0],
                        "kind": kind,
                        "rank": rank,
                        "nested_phrase": item["phrase"],
                        "count": int(item["count"]),
                        "total_count": "" if total is None else int(total),
                        "_order": order.get(region_name, 999),
                    }
                )
    rows.sort(key=lambda r: (r["phrase"], r["kind"], r["_order"], r["region_name"], r["rank"]))
    for row in rows:
        row.pop("_order")
    return rows


def write_csv(rows: list[dict], path: pathlib.Path) -> pathlib.Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    return path


def write_xlsx(cfg, rows: list[dict], path: pathlib.Path) -> pathlib.Path:
    """Лист на фразу: строки — вложенные запросы, колонки — регионы."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    order = region_order(cfg)
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    header_font = Font(bold=True)
    control_fill = PatternFill("solid", fgColor="EFEFEF")

    present = [p for p in cfg.phrases if any(r["phrase"] == p for r in rows)]
    for phrase in present:
        for kind, label in (("nested", ""), ("association", "ассоциации")):
            subset = [r for r in rows if r["phrase"] == phrase and r["kind"] == kind]
            if not subset:
                continue
            regions = sorted(
                {(r["region_name"], r["region_type"]) for r in subset},
                key=lambda item: (order.get(item[0], 999), item[0]),
            )
            totals: dict[str, int] = {}
            values: dict[str, dict[str, int]] = {}
            for row in subset:
                totals[row["nested_phrase"]] = totals.get(row["nested_phrase"], 0) + row["count"]
                values.setdefault(row["nested_phrase"], {})[row["region_name"]] = row["count"]
            ordered = sorted(totals, key=lambda p: (-totals[p], p))

            sheet = workbook.create_sheet(sheet_title(f"{phrase} {label}".strip(), "", False, used))
            sheet["A1"] = "Вложенный запрос" if kind == "nested" else "Ассоциация"
            sheet["B1"] = "всего"
            sheet["A1"].font = sheet["B1"].font = header_font
            for column, (name, rtype) in enumerate(regions, start=3):
                cell = sheet.cell(row=1, column=column, value=name)
                cell.font = header_font
                cell.alignment = Alignment(text_rotation=45, horizontal="left")
                if rtype != "campaign":
                    cell.fill = control_fill
            for line, nested in enumerate(ordered, start=2):
                sheet.cell(row=line, column=1, value=nested)
                sheet.cell(row=line, column=2, value=totals[nested])
                for column, (name, _) in enumerate(regions, start=3):
                    sheet.cell(row=line, column=column, value=values[nested].get(name))
            sheet.freeze_panes = "C2"
            sheet.column_dimensions["A"].width = 42
            sheet.column_dimensions["B"].width = 10
            for column in range(3, 3 + len(regions)):
                sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = 13

    if not workbook.sheetnames:
        raise SystemExit("нечего писать в xlsx: топ пуст")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build(cfg, raw_dir: pathlib.Path, out_dir: pathlib.Path, force: bool = False, log=print) -> dict:
    from .build import read_raw

    if out_dir.exists() and any(out_dir.iterdir()) and not force:
        raise SystemExit(
            f"{out_dir} уже заполнен — прогоны не перезаписываются. "
            "Пересобрать поверх: добавьте --force"
        )
    rows = to_rows(cfg, read_raw(raw_dir))
    if not rows:
        raise SystemExit("в сырых ответах нет вложенных запросов")
    csv_path = write_csv(rows, out_dir / "wordstat_top.csv")
    xlsx_path = write_xlsx(cfg, rows, out_dir / "wordstat_top.xlsx")
    summary = {
        "raw_dir": str(raw_dir),
        "rows": len(rows),
        "nested": sum(1 for r in rows if r["kind"] == "nested"),
        "associations": sum(1 for r in rows if r["kind"] == "association"),
        "phrases": len({r["phrase"] for r in rows}),
        "regions": len({r["region_name"] for r in rows}),
        "top_csv": str(csv_path),
        "top_xlsx": str(xlsx_path),
    }
    (out_dir / "build.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    log(f"строк: {len(rows)} (вложенных {summary['nested']}, ассоциаций {summary['associations']}); "
        f"регионов: {summary['regions']}")
    return summary
