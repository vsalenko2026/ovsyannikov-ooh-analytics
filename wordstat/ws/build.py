"""Сборка результата: сырые ответы -> wordstat_long.csv и wordstat_wide.xlsx."""

from __future__ import annotations

import csv
import datetime as dt
import json
import pathlib
import re

from . import periods

LONG_COLUMNS = [
    "date", "week_start", "week_end", "phrase", "region_name", "region_id",
    "region_type", "device", "count", "share", "is_partial",
]

INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def read_raw(directory: pathlib.Path) -> list[dict]:
    """Все сохранённые ответы прогона."""
    files = sorted(p for p in directory.glob("*.json") if p.name != "manifest.json")
    if not files:
        raise SystemExit(f"в {directory} нет сырых ответов")
    records = []
    for path in files:
        try:
            records.append({"path": path, **json.loads(path.read_text(encoding="utf-8"))})
        except ValueError as exc:
            raise SystemExit(f"битый файл {path}: {exc}") from exc
    return records


def resolve_anchor(cfg, records: list[dict]) -> str:
    """Якорь недели: из конфига или определённый по фактическим датам."""
    if cfg.week_anchor != "auto":
        return cfg.week_anchor
    dates = [
        item["date"]
        for record in records
        for item in (record["response"].get("results") or record["response"].get("dynamics") or [])
    ]
    if not dates:
        raise SystemExit("в ответах нет ни одной даты — определить границы недель невозможно")
    return periods.detect_week_anchor(dates)


def region_order(cfg) -> dict[str, int]:
    """Порядок колонок: кампания по числу экранов, затем контроль."""
    campaign = sorted(cfg.campaign_regions, key=lambda r: (-r.screens, r.name))
    control = sorted(cfg.control_regions, key=lambda r: r.name)
    return {r.name: i for i, r in enumerate([*campaign, *control])}


def to_rows(cfg, records: list[dict], anchor: str, today: dt.date | None = None) -> list[dict]:
    today = today or dt.date.today()
    by_id = {r.region_id: r for r in cfg.regions if r.region_id}
    order = region_order(cfg)
    rows = []
    for record in records:
        request = record["request"]
        meta = record.get("meta") or {}
        region_id = (request.get("regions") or [""])[0]
        region = by_id.get(str(region_id))
        region_name = meta.get("region_name") or (region.name if region else str(region_id))
        region_type = meta.get("region_type") or (region.type if region else "")
        device = (request.get("devices") or ["DEVICE_ALL"])[0]
        payload = record["response"].get("results") or record["response"].get("dynamics") or []
        for item in payload:
            day = periods.parse_api_date(item["date"])
            start, end = periods.period_bounds(day, request.get("period", cfg.period), anchor)
            share = item.get("share")
            rows.append(
                {
                    "date": day.isoformat(),
                    "week_start": start.isoformat(),
                    "week_end": end.isoformat(),
                    "phrase": request["phrase"],
                    "region_name": region_name,
                    "region_id": str(region_id),
                    "region_type": region_type,
                    "device": device,
                    "count": int(item["count"]),
                    "share": "" if share is None else float(share),
                    "is_partial": "0" if periods.is_closed(end, today) else "1",
                    "_order": order.get(region_name, 999),
                }
            )
    rows.sort(key=lambda r: (r["phrase"], r["device"], r["_order"], r["region_name"], r["date"]))
    for row in rows:
        row.pop("_order")
    return rows


def write_long(rows: list[dict], path: pathlib.Path) -> pathlib.Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=LONG_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    return path


def sheet_title(phrase: str, device: str, with_device: bool, used: set[str]) -> str:
    base = phrase if not with_device else f"{phrase} {device.replace('DEVICE_', '').lower()}"
    base = INVALID_SHEET_CHARS.sub(" ", base).strip()[:31] or "лист"
    title, n = base, 1
    while title in used:
        n += 1
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
    used.add(title)
    return title


def week_label(row: dict) -> str:
    start = periods.parse_date(row["week_start"])
    end = periods.parse_date(row["week_end"])
    return f"{start:%d.%m}–{end:%d.%m}"


def write_wide(cfg, rows: list[dict], path: pathlib.Path) -> pathlib.Path:
    """Широкий формат для глазной проверки: строки — недели, колонки — регионы."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    order = region_order(cfg)
    devices = sorted({r["device"] for r in rows})
    with_device = len(devices) > 1

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    header_font = Font(bold=True)
    control_fill = PatternFill("solid", fgColor="EFEFEF")
    partial_fill = PatternFill("solid", fgColor="FFF3E0")

    for phrase in cfg.phrases:
        for device in devices:
            subset = [r for r in rows if r["phrase"] == phrase and r["device"] == device]
            if not subset:
                continue
            regions = sorted(
                {(r["region_name"], r["region_type"]) for r in subset},
                key=lambda item: (order.get(item[0], 999), item[0]),
            )
            weeks: dict[str, dict] = {}
            for row in subset:
                week = weeks.setdefault(
                    row["week_start"],
                    {"label": week_label(row), "partial": row["is_partial"] == "1", "values": {}},
                )
                week["values"][row["region_name"]] = row["count"]

            sheet = workbook.create_sheet(sheet_title(phrase, device, with_device, used_titles))
            sheet["A1"] = "Неделя"
            sheet["B1"] = "закрыта"
            for column, (name, _) in enumerate(regions, start=3):
                cell = sheet.cell(row=1, column=column, value=name)
                cell.font = header_font
                cell.alignment = Alignment(text_rotation=45, horizontal="left")
            sheet["A2"] = "тип"
            for column, (_, rtype) in enumerate(regions, start=3):
                cell = sheet.cell(
                    row=2, column=column,
                    value="кампания" if rtype == "campaign" else "контроль",
                )
                if rtype != "campaign":
                    cell.fill = control_fill
            sheet["A1"].font = header_font
            sheet["B1"].font = header_font

            for line, week_start in enumerate(sorted(weeks), start=3):
                week = weeks[week_start]
                sheet.cell(row=line, column=1, value=week["label"])
                closed_cell = sheet.cell(row=line, column=2, value="нет" if week["partial"] else "да")
                if week["partial"]:
                    closed_cell.fill = partial_fill
                for column, (name, _) in enumerate(regions, start=3):
                    cell = sheet.cell(row=line, column=column, value=week["values"].get(name))
                    if week["partial"]:
                        cell.fill = partial_fill

            sheet.freeze_panes = "C3"
            sheet.column_dimensions["A"].width = 14
            sheet.column_dimensions["B"].width = 9
            for column in range(3, 3 + len(regions)):
                sheet.column_dimensions[sheet.cell(row=1, column=column).column_letter].width = 13

    if not workbook.sheetnames:
        raise SystemExit("нечего писать в xlsx: ни одной строки")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build(cfg, raw_dir: pathlib.Path, out_dir: pathlib.Path, force: bool = False,
          today: dt.date | None = None, log=print) -> dict:
    if out_dir.exists() and any(out_dir.iterdir()) and not force:
        raise SystemExit(
            f"{out_dir} уже заполнен — прогоны не перезаписываются. "
            "Пересобрать поверх: добавьте --force"
        )
    records = read_raw(raw_dir)
    anchor = resolve_anchor(cfg, records)
    rows = to_rows(cfg, records, anchor, today=today)
    if not rows:
        raise SystemExit("в сырых ответах нет данных")

    long_path = write_long(rows, out_dir / "wordstat_long.csv")
    wide_path = write_wide(cfg, rows, out_dir / "wordstat_wide.xlsx")

    partial = sorted({r["week_start"] for r in rows if r["is_partial"] == "1"})
    summary = {
        "raw_dir": str(raw_dir),
        "rows": len(rows),
        "week_anchor": anchor,
        "weeks": len({r["week_start"] for r in rows}),
        "regions": len({r["region_name"] for r in rows}),
        "phrases": len({r["phrase"] for r in rows}),
        "partial_weeks": partial,
        "long_csv": str(long_path),
        "wide_xlsx": str(wide_path),
    }
    (out_dir / "build.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    log(f"строк: {len(rows)}; недель: {summary['weeks']}; регионов: {summary['regions']}")
    log(f"дата недели в ответе API — {'начало' if anchor == 'start' else 'конец'} недели")
    if partial:
        log(f"незакрытые недели (is_partial=1): {', '.join(partial)}")
    return summary
