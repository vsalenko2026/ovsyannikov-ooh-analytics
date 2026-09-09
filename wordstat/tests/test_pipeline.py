"""Сквозная проверка без сети: выгрузка -> кэш -> сборка -> сверка."""

import csv
import datetime as dt
import functools
import pathlib
import sys
import tempfile
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from tests.fake_api import FakeSession, weekly_series
from ws import build as build_mod
from ws import config as config_mod
from ws import fetch as fetch_mod
from ws import reconcile as reconcile_mod
from ws.api import WordstatClient

CONFIG = """
api:
  base_url: https://searchapi.api.cloud.yandex.net/v2/wordstat
  pause_seconds: 0
  max_attempts: 3
  backoff_base_seconds: 0
  price_per_call_rub: 0.02
series:
  start_date: '2026-08-03'
  end_date: '2026-09-06'
  period: PERIOD_WEEKLY
  week_anchor: auto
  include_current_period: false
  devices: [DEVICE_ALL]
phrases:
  - овсянников мыло
  - овсянников косметика
regions_file: regions.yaml
paths:
  raw: raw
  output: output
"""

REGIONS = """
meta:
  level: oblast
regions:
- name: Ярославская область
  type: campaign
  region_id: '16'
  cities: [Ярославль, Рыбинск]
  screens: 7
  campaign_start: '2026-06-04'
- name: Тюменская область
  type: campaign
  region_id: '55'
  cities: [Тюмень]
  screens: 27
  campaign_start: '2026-07-17'
- name: Ростовская область
  type: control
  region_id: '39'
"""


class PipelineCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = pathlib.Path(self.tmp.name)
        (root / "config.yaml").write_text(CONFIG, encoding="utf-8")
        (root / "regions.yaml").write_text(REGIONS, encoding="utf-8")
        self.root = root
        self.cfg = config_mod.load(root / "config.yaml")
        self.session = FakeSession(generator=functools.partial(weekly_series, base=20, step=5))
        self.client = WordstatClient(
            api_key="k", folder_id="f", session=self.session,
            pause_seconds=0, backoff_base_seconds=0, sleep=lambda _s: None,
        )

    def tearDown(self):
        self.tmp.cleanup()

    def fetch(self, run_date=dt.date(2026, 9, 9)):
        return fetch_mod.fetch(self.cfg, self.client, run_date=run_date, log=lambda *_: None)


class TestFetch(PipelineCase):
    def test_plan_is_phrases_by_regions_by_devices(self):
        self.assertEqual(self.cfg.planned_calls(), 2 * 3 * 1)

    def test_one_region_per_call(self):
        self.fetch()
        for request in self.session.requests:
            self.assertEqual(len(request["body"]["regions"]), 1)

    def test_to_date_is_last_day_of_week(self):
        self.fetch()
        # API отвергает toDate, не совпадающий с концом периода
        self.assertTrue(all(r["body"]["toDate"].startswith("2026-09-06") for r in self.session.requests))
        self.assertTrue(all(r["body"]["fromDate"].startswith("2026-08-03") for r in self.session.requests))

    def test_second_run_is_fully_cached(self):
        first = self.fetch()
        calls_after_first = len(self.session.requests)
        second = self.fetch()
        self.assertEqual(first["fetched"], 6)
        self.assertEqual(second["fetched"], 0)
        self.assertEqual(second["from_cache"], 6)
        self.assertEqual(len(self.session.requests), calls_after_first)

    def test_cost_is_logged(self):
        summary = self.fetch()
        self.assertEqual(summary["billable_calls"], 6)
        self.assertAlmostEqual(summary["cost_rub"], 0.12)

    def test_raw_keeps_request_and_response(self):
        self.fetch()
        files = list((self.root / "raw" / "2026-09-09").glob("*.json"))
        self.assertEqual(len(files), 7)  # 6 ответов + manifest.json
        sample = next(p for p in files if p.name != "manifest.json")
        import json
        stored = json.loads(sample.read_text(encoding="utf-8"))
        self.assertIn("request", stored)
        self.assertIn("response", stored)
        self.assertIn("region_name", stored["meta"])


class TestBuild(PipelineCase):
    def build(self, today=dt.date(2026, 9, 9)):
        self.fetch()
        out = self.root / "output" / "2026-09-09"
        return build_mod.build(
            self.cfg, self.root / "raw" / "2026-09-09", out,
            today=today, log=lambda *_: None
        ), out

    def test_long_csv_columns_and_content(self):
        summary, out = self.build()
        with (out / "wordstat_long.csv").open(encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(list(rows[0]), build_mod.LONG_COLUMNS)
        # 2 фразы × 3 региона × 5 недель (03.08 … 06.09)
        self.assertEqual(len(rows), 2 * 3 * 5)
        self.assertEqual(summary["week_anchor"], "start")
        first = rows[0]
        self.assertEqual(first["week_start"], "2026-08-03")
        self.assertEqual(first["week_end"], "2026-08-09")
        self.assertEqual(first["is_partial"], "0")
        self.assertIn(first["region_type"], ("campaign", "control"))

    def test_partial_week_is_flagged(self):
        # Если считать «сегодня» 05.08, неделя 03–09.08 ещё не закрыта.
        summary, out = self.build(today=dt.date(2026, 8, 5))
        self.assertIn("2026-08-03", summary["partial_weeks"])
        with (out / "wordstat_long.csv").open(encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        partial = {r["week_start"] for r in rows if r["is_partial"] == "1"}
        self.assertEqual(partial, {"2026-08-03", "2026-08-10", "2026-08-17",
                                   "2026-08-24", "2026-08-31"})

    def test_wide_xlsx_has_sheet_per_phrase(self):
        _, out = self.build()
        from openpyxl import load_workbook
        book = load_workbook(out / "wordstat_wide.xlsx")
        self.assertEqual(len(book.sheetnames), 2)
        sheet = book[book.sheetnames[0]]
        self.assertEqual(sheet["A1"].value, "Неделя")
        # Кампания идёт первой и упорядочена по числу экранов
        self.assertEqual(sheet["C1"].value, "Тюменская область")
        self.assertEqual(sheet["E2"].value, "контроль")

    def test_previous_run_is_not_overwritten(self):
        _, out = self.build()
        with self.assertRaises(SystemExit):
            build_mod.build(self.cfg, self.root / "raw" / "2026-09-09", out, log=lambda *_: None)
        build_mod.build(self.cfg, self.root / "raw" / "2026-09-09", out, force=True, log=lambda *_: None)

    def test_mixed_week_grid_stops_the_build(self):
        # Ответ с датами не на одной границе недели — сборка обязана упасть,
        # а не достроить границы наугад.
        records = [{
            "request": {"phrase": "овсянников мыло", "regions": ["16"], "period": "PERIOD_WEEKLY"},
            "meta": {},
            "response": {"results": [{"date": "2026-08-03", "count": "1"},
                                     {"date": "2026-08-06", "count": "2"}]},
        }]
        with self.assertRaises(ValueError):
            build_mod.resolve_anchor(self.cfg, records)


class TestReconcile(PipelineCase):
    def setUp(self):
        super().setUp()
        self.fetch()
        self.out = self.root / "output" / "2026-09-09"
        build_mod.build(self.cfg, self.root / "raw" / "2026-09-09", self.out,
                        today=dt.date(2026, 9, 9), log=lambda *_: None)
        self.api = reconcile_mod.load_api(self.out / "wordstat_long.csv")

    def manual_file(self, scale=1.0, region_name="Ярославская область"):
        path = self.root / "manual.csv"
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["фраза", "регион", "дата", "частота"])
            for row in self.api:
                if row["region"] == "Ярославская область" and row["phrase"] == "овсянников мыло":
                    writer.writerow([row["phrase"], region_name,
                                     row["week_start"].isoformat(), round(row["count"] * scale)])
        return path

    def test_identical_data_passes(self):
        rows = reconcile_mod.compare(self.cfg, reconcile_mod.load_manual(self.manual_file()), self.api)
        summary = reconcile_mod.verdict(rows)
        self.assertTrue(summary["ok"], summary["stop_reasons"])
        self.assertEqual(summary["missing"], 0)

    def test_systematic_shift_stops(self):
        rows = reconcile_mod.compare(self.cfg, reconcile_mod.load_manual(self.manual_file(scale=1.3)), self.api)
        summary = reconcile_mod.verdict(rows)
        self.assertFalse(summary["ok"])
        self.assertTrue(any("отношение" in r or "расхождение" in r for r in summary["stop_reasons"]))

    def test_city_name_maps_to_region(self):
        # В ручной выгрузке регион может быть назван городом
        rows = reconcile_mod.compare(
            self.cfg, reconcile_mod.load_manual(self.manual_file(region_name="Ярославль")), self.api
        )
        self.assertTrue(all(r["api"] is not None for r in rows))

    def test_markdown_report_renders(self):
        rows = reconcile_mod.compare(self.cfg, reconcile_mod.load_manual(self.manual_file()), self.api)
        text = reconcile_mod.render_markdown(rows, reconcile_mod.verdict(rows))
        self.assertIn("| Фраза | Регион | Неделя | Ручная | API | Расхождение |", text)
        self.assertIn("сверка пройдена", text)


if __name__ == "__main__":
    unittest.main()
