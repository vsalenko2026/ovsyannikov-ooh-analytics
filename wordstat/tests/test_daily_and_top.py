"""Дневная выгрузка и снимок топа вложенных запросов — без сети."""

import csv
import dataclasses
import datetime as dt
import pathlib
import sys
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from tests.fake_api import FakeSession, dispatch
from tests.test_pipeline import PipelineCase
from ws import build as build_mod
from ws import fetch as fetch_mod
from ws import periods
from ws import top as top_mod


class DailyCase(PipelineCase):
    def setUp(self):
        super().setUp()
        self.session = FakeSession(generator=dispatch)
        self.client = self.client.__class__(
            api_key="k", folder_id="f", session=self.session,
            pause_seconds=0, backoff_base_seconds=0, sleep=lambda _s: None,
        )
        # конфиг тестовый: ряд с 03.08, «сегодня» — 09.09.2026
        self.daily = dataclasses.replace(self.cfg, period=periods.PERIOD_DAILY)


class TestDaily(DailyCase):
    def test_start_is_clamped_to_lookback(self):
        cfg = dataclasses.replace(self.daily, start_date=dt.date(2026, 3, 30))
        start, clamped = cfg.effective_start_date(dt.date(2026, 9, 9))
        self.assertTrue(clamped)
        self.assertEqual(start, dt.date(2026, 7, 12))    # 60 дней назад включительно

    def test_start_is_kept_when_inside_lookback(self):
        start, clamped = self.daily.effective_start_date(dt.date(2026, 9, 9))
        self.assertFalse(clamped)
        self.assertEqual(start, dt.date(2026, 8, 3))

    def test_end_is_yesterday(self):
        # текущий день не закрыт и по умолчанию не запрашивается
        cfg = dataclasses.replace(self.daily, end_date=None)
        self.assertEqual(cfg.effective_end_date(dt.date(2026, 9, 9)), dt.date(2026, 9, 8))

    def test_run_dir_does_not_collide_with_weekly(self):
        weekly = fetch_mod.run_dir(self.cfg, dt.date(2026, 9, 9)).name
        daily = fetch_mod.run_dir(self.daily, dt.date(2026, 9, 9)).name
        self.assertEqual(weekly, "2026-09-09")
        self.assertEqual(daily, "2026-09-09-day")
        self.assertNotEqual(weekly, daily)

    def test_daily_build_gives_one_row_per_day(self):
        fetch_mod.fetch(self.daily, self.client, run_date=dt.date(2026, 9, 9), log=lambda *_: None)
        raw = fetch_mod.run_dir(self.daily, dt.date(2026, 9, 9))
        out = self.root / "output" / raw.name
        summary = build_mod.build(self.daily, raw, out, today=dt.date(2026, 9, 9), log=lambda *_: None)
        self.assertEqual(summary["period"], periods.PERIOD_DAILY)
        with (out / "wordstat_long.csv").open(encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        days = (dt.date(2026, 9, 6) - dt.date(2026, 8, 3)).days + 1
        self.assertEqual(len(rows), 2 * 3 * days)
        # у дня границы периода совпадают с самим днём
        self.assertTrue(all(r["week_start"] == r["date"] == r["week_end"] for r in rows))

    def test_mixed_periods_in_one_dir_are_refused(self):
        records = [
            {"request": {"phrase": "а", "regions": ["16"], "period": "PERIOD_WEEKLY"},
             "meta": {}, "response": {"results": []}},
            {"request": {"phrase": "а", "regions": ["16"], "period": "PERIOD_DAILY"},
             "meta": {}, "response": {"results": []}},
        ]
        with self.assertRaises(SystemExit):
            build_mod.records_period(self.cfg, records)


class TestTop(DailyCase):
    def export(self):
        run = dt.date(2026, 9, 9)
        top_mod.fetch(self.cfg, self.client, run_date=run, limit=5, log=lambda *_: None)
        raw = top_mod.run_dir(self.cfg, run)
        out = self.root / "output" / raw.name
        return top_mod.build(self.cfg, raw, out, log=lambda *_: None), out

    def test_covers_every_phrase_and_region(self):
        summary, _ = self.export()
        self.assertEqual(summary["phrases"], 2)
        self.assertEqual(summary["regions"], 3)
        # 3 вложенных + 2 ассоциации на каждую пару
        self.assertEqual(summary["nested"], 2 * 3 * 3)
        self.assertEqual(summary["associations"], 2 * 3 * 2)

    def test_csv_columns(self):
        _, out = self.export()
        with (out / "wordstat_top.csv").open(encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(list(rows[0]), top_mod.COLUMNS)
        self.assertIn(rows[0]["kind"], ("nested", "association"))
        self.assertEqual(rows[0]["fetched_at"][:2], "20")

    def test_xlsx_sheet_per_phrase_and_kind(self):
        _, out = self.export()
        from openpyxl import load_workbook
        book = load_workbook(out / "wordstat_top.xlsx")
        self.assertEqual(len(book.sheetnames), 4)   # 2 фразы × (вложенные + ассоциации)
        sheet = book[book.sheetnames[0]]
        self.assertEqual(sheet["A1"].value, "Вложенный запрос")
        self.assertEqual(sheet["B1"].value, "всего")
        self.assertEqual(sheet["C1"].value, "Тюменская область")

    def test_second_run_is_cached(self):
        self.export()
        calls = len(self.session.requests)
        summary = top_mod.fetch(self.cfg, self.client, run_date=dt.date(2026, 9, 9),
                                limit=5, log=lambda *_: None)
        self.assertEqual(summary["fetched"], 0)
        self.assertEqual(summary["from_cache"], 6)
        self.assertEqual(len(self.session.requests), calls)

    def test_top_does_not_collide_with_dynamics_dir(self):
        self.assertEqual(top_mod.run_dir(self.cfg, dt.date(2026, 9, 9)).name, "2026-09-09-top")


if __name__ == "__main__":
    unittest.main()
